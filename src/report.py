# Компактный Excel-отчёт: одна книга, несколько листов.
# Собирается из сохранённых файлов l2_*.csv и l2_selection.json,
# обучение не выполняется.

import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.visualisation import MODEL_LABELS, russian_feature_name

SHEET_PARAMETERS = "Параметры"
SHEET_CV = "Сводка CV"
SHEET_COEFFICIENTS = "Коэффициенты"
SHEET_TEST = "Тест"
SHEET_CMP_CV = "Сравнение CV"
SHEET_CMP_TEST = "Сравнение тест"
SHEET_CMP_RF = "Лес, сетка"
SHEET_CMP_SOLVERS = "Решатели"

HEADER_FILL = PatternFill("solid", fgColor="E4E3DF")
SELECTED_FILL = PatternFill("solid", fgColor="FBD9CB")

# Параметры выбора: ключ в selection.json, подпись, описание
PARAMETER_ROWS = [
    ("C", "C (обратная сила регуляризации)", "выбрано по правилу одной стандартной ошибки"),
    ("l2", "l2 (коэффициент штрафа)", "l2 = 1 / (C · n_dev)"),
    ("threshold", "Порог классификации", "выбран по out-of-fold вероятностям"),
    ("min_sensitivity", "Требуемая чувствительность", "ограничение при выборе порога"),
    ("oof_sensitivity", "Чувствительность при пороге (OOF)", "на development, out-of-fold"),
    ("oof_specificity", "Специфичность при пороге (OOF)", "на development, out-of-fold"),
    ("n_dev", "Объектов в development", ""),
    ("n_test", "Объектов в тесте", ""),
    ("n_splits", "Число фолдов", ""),
    ("n_repeats", "Число повторов CV", ""),
    ("learning_rate", "Скорость обучения", ""),
    ("epochs_cap", "Лимит эпох", ""),
    ("tol", "Порог остановки по норме градиента", ""),
    ("cv_converged_share", "Доля сошедшихся обучений в CV", ""),
    ("final_epochs", "Эпох у итоговой модели", ""),
    ("final_converged", "Итоговая модель сошлась", ""),
    ("random_state", "random_state", ""),
    ("test_evaluated", "Оценка на тесте выполнена", ""),
    ("elapsed_seconds", "Время прогона, с", ""),
]

CV_COLUMNS = {
    "C": "C",
    "l2": "l2",
    "log_loss_mean": "log-loss, среднее",
    "log_loss_std": "log-loss, σ по фолдам",
    "brier_mean": "Brier, среднее",
    "brier_std": "Brier, σ по фолдам",
    "roc_auc_mean": "ROC-AUC, среднее",
    "roc_auc_std": "ROC-AUC, σ по фолдам",
    "converged_share": "Доля сошедшихся",
    "epochs_mean": "Эпох в среднем",
}

METRIC_NAMES = {
    "accuracy": "Точность (accuracy)",
    "sensitivity": "Чувствительность",
    "specificity": "Специфичность",
    "precision": "Precision",
    "f1": "F1",
    "roc_auc": "ROC-AUC",
    "auprc": "AUPRC",
    "log_loss": "log-loss",
    "brier": "Brier",
}

CI_METHODS = {
    "Clopper-Pearson": "Клоппер–Пирсон",
    "bootstrap": "бутстреп (2000)",
}


def parameters_table(selection):
    rows = []
    for key, label, note in PARAMETER_ROWS:
        if key in selection:
            rows.append({"Параметр": label, "Значение": selection[key],
                         "Комментарий": note})
    return pd.DataFrame(rows)


def cv_table(summary):
    # По возрастанию C: слева сильная регуляризация, справа слабая
    table = summary.sort_values("C").reset_index(drop=True)
    return table[list(CV_COLUMNS)].rename(columns=CV_COLUMNS)


def coefficients_table(coefficients):
    table = coefficients.copy()
    table.insert(1, "Признак (рус.)",
                 table["feature"].map(russian_feature_name))
    return table.rename(columns={
        "feature": "Признак",
        "coef_standardized": "Коэффициент (стандартизованные признаки)",
        "coef_original_units": "Коэффициент (исходные единицы)",
    })


def test_table(metrics):
    table = pd.DataFrame({
        "Метрика": metrics["metric"].map(METRIC_NAMES).fillna(metrics["metric"]),
        "Значение": metrics["value"],
        "ДИ 95%, нижняя": metrics["ci_low"],
        "ДИ 95%, верхняя": metrics["ci_high"],
        "Метод ДИ": metrics["ci_method"].map(CI_METHODS).fillna(metrics["ci_method"]),
    })
    return table


def cmp_cv_table(cmp_summary):
    table = cmp_summary.copy()
    table["model"] = table["model"].map(MODEL_LABELS).fillna(table["model"])
    return table.rename(columns={
        "model": "Модель",
        "config": "Параметры",
        "log_loss_mean": "log-loss, среднее",
        "log_loss_std": "log-loss, σ по фолдам",
        "brier_mean": "Brier, среднее",
        "brier_std": "Brier, σ по фолдам",
        "roc_auc_mean": "ROC-AUC, среднее",
        "roc_auc_std": "ROC-AUC, σ по фолдам",
        "threshold": "Порог",
        "oof_sensitivity": "Чувствительность (OOF)",
        "oof_specificity": "Специфичность (OOF)",
    })


def _comma(value):
    return f"{value:.4f}".replace(".", ",")


# Одна строка на метрику; для каждой модели пара столбцов «значение» и «ДИ».
# Имена столбцов вида «модель|подпись»: первая часть станет объединённой шапкой
def cmp_test_wide(cmp_metrics):
    present = set(cmp_metrics["model"])
    models = [model for model in MODEL_LABELS if model in present]

    rows = []
    for metric, label in METRIC_NAMES.items():
        part = cmp_metrics[cmp_metrics["metric"] == metric]
        if part.empty:
            continue

        row = {"Метрика": label}
        for model in models:
            line = part[part["model"] == model].iloc[0]
            row[f"{model}|Значение"] = float(line["value"])
            row[f"{model}|95% ДИ"] = (
                f"[{_comma(line['ci_low'])}; {_comma(line['ci_high'])}]"
            )

        method = part.iloc[0]["ci_method"]
        row["Метод ДИ"] = CI_METHODS.get(method, method)
        rows.append(row)

    return pd.DataFrame(rows)


# Запись таблицы с двухуровневой шапкой (модель над парой столбцов)
def _write_two_level(writer, name, table):
    models = []
    for column in table.columns:
        if "|" in column and column.split("|")[0] not in models:
            models.append(column.split("|")[0])

    flat = table.rename(columns=lambda c: c.split("|")[1] if "|" in c else c)
    flat.to_excel(writer, sheet_name=name, index=False, startrow=1)
    sheet = writer.sheets[name]

    sheet.cell(1, 1, table.columns[0])
    sheet.merge_cells(start_row=1, end_row=2, start_column=1, end_column=1)

    column = 2
    for model in models:
        sheet.cell(1, column, MODEL_LABELS[model])
        sheet.merge_cells(start_row=1, end_row=1,
                          start_column=column, end_column=column + 1)
        column += 2

    sheet.cell(1, column, table.columns[-1])
    sheet.merge_cells(start_row=1, end_row=2, start_column=column,
                      end_column=column)
    return sheet


def cmp_rf_table(rf_grid):
    return rf_grid.rename(columns={
        "min_samples_leaf": "min_samples_leaf",
        "log_loss_mean": "log-loss, среднее",
        "log_loss_std": "log-loss, σ по фолдам",
        "brier_mean": "Brier, среднее",
        "brier_std": "Brier, σ по фолдам",
        "roc_auc_mean": "ROC-AUC, среднее",
        "roc_auc_std": "ROC-AUC, σ по фолдам",
    })


def cmp_solver_table(solvers):
    return solvers.rename(columns={
        "solver": "Решатель sklearn",
        "n_iter": "Итераций",
        "max_abs_diff_weights": "Макс. |Δ весов| против собственной модели",
        "abs_diff_intercept": "|Δ свободного члена|",
    })


def _format_sheet(worksheet, number_format="0.0000", selected_row=None,
                  header_rows=1):
    for header_row in range(1, header_rows + 1):
        for cell in worksheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(
                wrap_text=True, vertical="center", horizontal="center"
                if header_rows > 1 else None,
            )
    worksheet.freeze_panes = f"A{header_rows + 1}" if header_rows == 1 else "B3"

    # Ширина по данным; текст объединённой шапки в расчёт не входит
    for index, column in enumerate(worksheet.columns, start=1):
        longest = max(
            (len(str(cell.value)) for cell in column[header_rows:]
             if cell.value is not None),
            default=0,
        )
        title = len(str(column[header_rows - 1].value or ""))
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(12, longest + 2, title + 2 if header_rows == 1 else 0), 48
        )

    for row in worksheet.iter_rows(min_row=header_rows + 1):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = number_format
            if header_rows > 1 and cell.column > 1:
                cell.alignment = Alignment(horizontal="center")

    if selected_row is not None:
        for cell in worksheet[selected_row]:
            cell.fill = SELECTED_FILL


def build_report(results_dir, output_path=None, prefix="l2_", cmp_prefix="cmp_"):
    # Возвращает путь к книге и список включённых листов
    results_dir = Path(results_dir)
    output_path = Path(output_path) if output_path else results_dir / f"{prefix}report.xlsx"

    def path(name):
        return results_dir / f"{prefix}{name}"

    with open(path("selection.json"), encoding="utf-8") as file:
        selection = json.load(file)

    summary = pd.read_csv(path("cv_summary.csv"))
    coefficients = pd.read_csv(path("coefficients.csv"))

    cv = cv_table(summary)
    selected_index = int(np.argmin(np.abs(cv["C"].to_numpy() - selection["C"])))

    sheets = [
        (SHEET_PARAMETERS, parameters_table(selection), "0.0000", None),
        (SHEET_CV, cv, "0.0000", selected_index + 2),
        (SHEET_COEFFICIENTS, coefficients_table(coefficients), "0.0000", None),
    ]

    # Лист с тестом появляется только после финального прогона
    if path("test_metrics.csv").exists():
        metrics = pd.read_csv(path("test_metrics.csv"))
        sheets.append((SHEET_TEST, test_table(metrics), "0.0000", None))

    # Сравнение моделей: листы появляются, если запускался compare_models.py
    def cmp_path(name):
        return results_dir / f"{cmp_prefix}{name}"

    if cmp_path("cv_summary.csv").exists():
        sheets.append((SHEET_CMP_CV, cmp_cv_table(
            pd.read_csv(cmp_path("cv_summary.csv"))), "0.0000", None))
    if cmp_path("test_metrics.csv").exists():
        sheets.append((SHEET_CMP_TEST, cmp_test_wide(
            pd.read_csv(cmp_path("test_metrics.csv"))), "0.0000", None, 2))
    if cmp_path("rf_grid.csv").exists():
        sheets.append((SHEET_CMP_RF, cmp_rf_table(
            pd.read_csv(cmp_path("rf_grid.csv"))), "0.0000", None))
    if cmp_path("solver_check.csv").exists():
        sheets.append((SHEET_CMP_SOLVERS, cmp_solver_table(
            pd.read_csv(cmp_path("solver_check.csv"))), "0.00E+00", None))

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for name, table, number_format, selected_row, *rest in sheets:
            header_rows = rest[0] if rest else 1
            if header_rows == 2:
                sheet = _write_two_level(writer, name, table)
            else:
                table.to_excel(writer, sheet_name=name, index=False)
                sheet = writer.sheets[name]
            _format_sheet(sheet, number_format, selected_row, header_rows)

    return output_path, [name for name, *_ in sheets]
