import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import make_report
from src import report

PREFIX = make_report.PREFIX


def write_results(folder, with_test):
    c = np.logspace(-2, 0.5, 5)
    pd.DataFrame({
        "C": c[::-1],  # намеренно в обратном порядке
        "l2": 1 / (c[::-1] * 100),
        "log_loss_mean": 0.1, "log_loss_std": 0.02,
        "brier_mean": 0.03, "brier_std": 0.01,
        "roc_auc_mean": 0.99, "roc_auc_std": 0.005,
        "converged_share": 1.0, "epochs_mean": 1000.0,
    }).to_csv(folder / f"{PREFIX}cv_summary.csv", index=False)

    pd.DataFrame({
        "feature": ["(intercept)", "worst texture", "mean radius"],
        "coef_standardized": [-0.4, 0.7, -0.5],
        "coef_original_units": [-20.0, 0.1, -0.2],
    }).to_csv(folder / f"{PREFIX}coefficients.csv", index=False)

    with open(folder / f"{PREFIX}selection.json", "w", encoding="utf-8") as f:
        json.dump({"C": float(c[2]), "l2": 0.01, "threshold": 0.44,
                   "n_dev": 100, "test_evaluated": with_test}, f)

    if with_test:
        pd.DataFrame({
            "metric": ["sensitivity", "roc_auc"],
            "value": [0.95, 0.99],
            "ci_low": [0.85, 0.97],
            "ci_high": [0.99, 1.0],
            "ci_method": ["Clopper-Pearson", "bootstrap"],
        }).to_csv(folder / f"{PREFIX}test_metrics.csv", index=False)


def test_parameters_table_skips_missing_keys():
    table = report.parameters_table({"C": 0.18, "threshold": 0.44})

    assert len(table) == 2
    assert list(table.columns) == ["Параметр", "Значение", "Комментарий"]


def test_cv_table_is_sorted_by_c_and_renamed():
    summary = pd.DataFrame({
        "C": [1.0, 0.1], "l2": [0.01, 0.1],
        "log_loss_mean": [0.2, 0.1], "log_loss_std": [0.0, 0.0],
        "brier_mean": [0.0, 0.0], "brier_std": [0.0, 0.0],
        "roc_auc_mean": [0.9, 0.9], "roc_auc_std": [0.0, 0.0],
        "converged_share": [1.0, 1.0], "epochs_mean": [10.0, 10.0],
    })

    table = report.cv_table(summary)

    assert list(table["C"]) == [0.1, 1.0]
    assert "log-loss, среднее" in table.columns


def test_coefficients_table_adds_russian_names():
    coefficients = pd.DataFrame({
        "feature": ["(intercept)", "worst texture"],
        "coef_standardized": [0.1, 0.2],
        "coef_original_units": [0.3, 0.4],
    })

    table = report.coefficients_table(coefficients)

    assert table.loc[1, "Признак (рус.)"] == "текстура (наихудшее)"


def test_build_report_without_test_has_three_sheets(tmp_path):
    write_results(tmp_path, with_test=False)

    path, sheets = report.build_report(tmp_path, prefix=PREFIX)

    assert path.exists()
    assert sheets == [report.SHEET_PARAMETERS, report.SHEET_CV,
                      report.SHEET_COEFFICIENTS]
    assert load_workbook(path).sheetnames == sheets


def test_build_report_with_test_adds_test_sheet(tmp_path):
    write_results(tmp_path, with_test=True)

    path, sheets = report.build_report(tmp_path, prefix=PREFIX)

    assert sheets[-1] == report.SHEET_TEST
    sheet = load_workbook(path)[report.SHEET_TEST]
    assert sheet.max_row == 3  # заголовок и две метрики
    assert sheet["A2"].value == "Чувствительность"
    assert sheet["E2"].value == "Клоппер–Пирсон"


def test_selected_c_row_is_highlighted(tmp_path):
    write_results(tmp_path, with_test=False)

    path, _ = report.build_report(tmp_path, prefix=PREFIX)
    sheet = load_workbook(path)[report.SHEET_CV]

    filled = [row for row in range(2, sheet.max_row + 1)
              if sheet.cell(row, 1).fill.fgColor.rgb.endswith("FBD9CB")]
    assert filled == [4]  # третья по возрастанию C строка данных


def test_make_report_main_returns_path(tmp_path):
    write_results(tmp_path, with_test=False)

    path = make_report.main(tmp_path)

    assert path == tmp_path / f"{PREFIX}report.xlsx"
    assert path.exists()
