import json

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

import compare_models as cmp
import make_figures
from src import report
from src import visualisation as viz
from src.models import fit


# ---------- Вспомогательные данные ----------

def make_small_dataset(n=160, seed=0):
    rng = np.random.default_rng(seed)
    X = rng.normal(size=(n, 3))
    logits = 1.2 * X[:, 0] - 0.8 * X[:, 1]
    y = (rng.random(n) < 1 / (1 + np.exp(-logits))).astype(int)
    return X, y


def small_forest(leaf, random_state=0, n_train=None):
    return cmp.make_forest(leaf, random_state, n_train, n_estimators=20)


def make_scores(n=200, seed=0):
    rng = np.random.default_rng(seed)
    y = (rng.random(n) < 0.4).astype(int)
    prob = np.clip(0.4 * y + rng.normal(0.3, 0.2, n), 0.01, 0.99)
    return y, prob


# ---------- Эквивалентность наших весов и sklearn ----------

def test_sklearn_with_scaled_c_matches_own_fit():
    X, y = make_small_dataset()
    X = (X - X.mean(axis=0)) / X.std(axis=0)
    l2 = 0.02

    weights, bias = fit(X, y, learning_rate=0.5, epochs=50000, l2=l2, tol=1e-10)
    model = cmp.make_sklearn_lr(l2, n_train=len(y)).fit(X, y)

    assert np.max(np.abs(weights - model.coef_[0])) < 1e-4
    assert abs(bias - model.intercept_[0]) < 1e-4


def test_sklearn_c_depends_on_training_size():
    small = cmp.make_sklearn_lr(0.02, n_train=100)
    large = cmp.make_sklearn_lr(0.02, n_train=400)

    # Один и тот же l2 при большей выборке даёт меньший C
    assert small.C == pytest.approx(0.5)
    assert large.C == pytest.approx(0.125)


def test_solver_check_reports_small_differences():
    X, y = make_small_dataset()
    X_scaled = (X - X.mean(axis=0)) / X.std(axis=0)
    l2 = 0.02
    weights, bias = fit(X_scaled, y, learning_rate=0.5, epochs=50000,
                        l2=l2, tol=1e-10)
    names = ["a", "b", "c"]
    own = pd.DataFrame({
        "feature": ["(intercept)"] + names,
        "coef_standardized": [bias, *weights],
    })

    table = cmp.solver_check(X, y, l2, own, names)

    assert list(table["solver"]) == ["lbfgs", "newton-cholesky"]
    assert (table["max_abs_diff_weights"] < 1e-3).all()


# ---------- CV и выбор параметра ----------

def test_cross_validate_model_shapes_and_range():
    X, y = make_small_dataset()

    table, oof = cmp.cross_validate_model(
        X, y, small_forest, {1: 1, 5: 5}, scale=False,
        n_splits=3, n_repeats=2, desc="тест",
    )

    assert len(table) == 3 * 2 * 2
    assert set(table["config"]) == {1, 5}
    for leaf in (1, 5):
        assert oof[leaf].shape == (2, len(y))
        assert np.all((oof[leaf] >= 0) & (oof[leaf] <= 1))


def test_cross_validate_model_uses_same_folds_as_own_cv():
    from sklearn.model_selection import RepeatedStratifiedKFold

    X, y = make_small_dataset()
    seen = []

    def spy(params, random_state, n_train):
        seen.append(n_train)
        return cmp.make_sklearn_lr(params, random_state, n_train)

    cmp.cross_validate_model(
        X, y, spy, {0.02: 0.02}, scale=True,
        n_splits=4, n_repeats=1, random_state=7, desc="тест",
    )

    expected = [
        len(train) for train, _ in
        RepeatedStratifiedKFold(n_splits=4, n_repeats=1, random_state=7).split(X, y)
    ]
    assert seen == expected


def test_select_best_picks_minimum_log_loss():
    summary = pd.DataFrame({
        "config": [1, 3, 8],
        "log_loss_mean": [0.14, 0.13, 0.15],
    })

    assert cmp.select_best(summary) == 3


def test_summarize_configs_mean_and_std():
    fold_table = pd.DataFrame({
        "config": [1, 1],
        "log_loss": [0.2, 0.4],
        "brier": [0.1, 0.1],
        "roc_auc": [0.9, 0.9],
    })

    summary = cmp.summarize_configs(fold_table).iloc[0]

    assert summary["log_loss_mean"] == pytest.approx(0.3)
    assert summary["log_loss_std"] == pytest.approx(np.sqrt(0.02))


# ---------- main ----------

def test_main_requires_results_of_main_py(tmp_path):
    with pytest.raises(FileNotFoundError, match="main.py"):
        cmp.main(results_dir=tmp_path)


# ---------- Рисунок и отчёт ----------

def test_plot_model_comparison_creates_files(tmp_path):
    y, prob = make_scores()
    probs = {
        "own_lr": prob,
        "sklearn_lr": prob,
        "random_forest": np.clip(prob + 0.05, 0, 1),
    }

    paths = viz.plot_model_comparison(y, probs, tmp_path / "fig7", n_bins=8)

    assert len(paths) == 2
    assert all(p.exists() and p.stat().st_size > 1000 for p in paths)


def write_comparison_results(folder, with_test):
    prefix = cmp.PREFIX
    y, prob = make_scores(120)

    c = np.logspace(-2, 0.5, 5)
    pd.DataFrame({
        "C": c, "l2": 1 / (c * 100),
        "log_loss_mean": 0.1, "log_loss_std": 0.02,
        "brier_mean": 0.03, "brier_std": 0.01,
        "roc_auc_mean": 0.99, "roc_auc_std": 0.005,
        "converged_share": 1.0, "epochs_mean": 1000.0,
    }).to_csv(folder / f"{prefix}cv_summary.csv", index=False)
    pd.DataFrame({
        "feature": ["(intercept)", "worst texture"],
        "coef_standardized": [-0.4, 0.7],
        "coef_original_units": [-20.0, 0.1],
    }).to_csv(folder / f"{prefix}coefficients.csv", index=False)
    pd.DataFrame({"loss": np.linspace(0.7, 0.1, 50),
                  "grad_norm": np.logspace(0, -6, 50)}
                 ).to_csv(folder / f"{prefix}final_history.csv", index=False)
    with open(folder / f"{prefix}selection.json", "w", encoding="utf-8") as f:
        json.dump({"C": float(c[2]), "l2": 0.01, "threshold": 0.44,
                   "min_sensitivity": 0.95, "n_splits": 5, "tol": 1e-6,
                   "final_epochs": 50, "learning_rate": 0.5}, f)

    own_oof = pd.DataFrame({"y_true": y})
    cmp_oof = pd.DataFrame({"y_true": y})
    for r in range(2):
        own_oof[f"y_prob_oof_r{r}"] = prob
        cmp_oof[f"y_prob_sklearn_lr_r{r}"] = prob
        cmp_oof[f"y_prob_random_forest_r{r}"] = np.clip(prob + 0.05, 0, 1)
    own_oof.to_csv(folder / f"{prefix}oof_predictions.csv")
    cmp_oof.to_csv(folder / f"{cmp.CMP_PREFIX}oof_predictions.csv")

    pd.DataFrame({
        "model": list(cmp.MODEL_ORDER),
        "config": ["C = 0.18", "C = 0.18", "500 деревьев"],
        "log_loss_mean": 0.1, "log_loss_std": 0.02,
        "brier_mean": 0.03, "brier_std": 0.01,
        "roc_auc_mean": 0.99, "roc_auc_std": 0.005,
        "threshold": 0.44, "oof_sensitivity": 0.95, "oof_specificity": 0.98,
    }).to_csv(folder / f"{cmp.CMP_PREFIX}cv_summary.csv", index=False)
    pd.DataFrame({"min_samples_leaf": [1, 3], "log_loss_mean": 0.14,
                  "log_loss_std": 0.03, "brier_mean": 0.03, "brier_std": 0.01,
                  "roc_auc_mean": 0.98, "roc_auc_std": 0.01}
                 ).to_csv(folder / f"{cmp.CMP_PREFIX}rf_grid.csv", index=False)
    pd.DataFrame({"solver": ["lbfgs"], "n_iter": [29],
                  "max_abs_diff_weights": [3e-5], "abs_diff_intercept": [1e-5]}
                 ).to_csv(folder / f"{cmp.CMP_PREFIX}solver_check.csv", index=False)

    if with_test:
        pd.DataFrame({
            "model": ["own_lr", "sklearn_lr", "random_forest"],
            "metric": "roc_auc", "value": 0.99, "ci_low": 0.97,
            "ci_high": 1.0, "ci_method": "bootstrap",
        }).to_csv(folder / f"{cmp.CMP_PREFIX}test_metrics.csv", index=False)
        pd.DataFrame({
            "y_true": y[:60], "y_prob_own_lr": prob[:60],
            "y_prob_sklearn_lr": prob[:60], "y_prob_random_forest": prob[:60],
        }).to_csv(folder / f"{cmp.CMP_PREFIX}test_predictions.csv")


def test_report_adds_comparison_sheets(tmp_path):
    write_comparison_results(tmp_path, with_test=True)

    path, sheets = report.build_report(tmp_path, prefix=cmp.PREFIX)

    for name in (report.SHEET_CMP_CV, report.SHEET_CMP_TEST,
                 report.SHEET_CMP_RF, report.SHEET_CMP_SOLVERS):
        assert name in sheets
    sheet = load_workbook(path)[report.SHEET_CMP_TEST]
    # Двухуровневая шапка: модели над парами «значение / 95% ДИ»
    assert sheet["B1"].value == viz.MODEL_LABELS["own_lr"]
    assert sheet["B2"].value == "Значение"
    assert sheet["C2"].value == "95% ДИ"
    # Одна строка на метрику
    assert sheet["A3"].value == "ROC-AUC"
    assert sheet["C3"].value.startswith("[")
    assert sheet.max_row == 3


def test_report_without_comparison_has_no_comparison_sheets(tmp_path):
    write_comparison_results(tmp_path, with_test=False)
    for name in ("cv_summary.csv", "rf_grid.csv", "solver_check.csv",
                 "oof_predictions.csv"):
        (tmp_path / f"{cmp.CMP_PREFIX}{name}").unlink()

    _, sheets = report.build_report(tmp_path, prefix=cmp.PREFIX)

    assert not any(name.startswith(("Сравнение", "Лес", "Решатели"))
                   for name in sheets)


def test_make_figures_builds_comparison_figures(tmp_path):
    write_comparison_results(tmp_path, with_test=True)

    make_figures.main(tmp_path)

    names = {p.name for p in (tmp_path / "figures").iterdir()}
    assert "fig7_comparison_oof.png" in names
    assert "fig7_comparison_test.pdf" in names
